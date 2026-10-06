"""PART B1: Reports"""
import os, subprocess

FILES = {}
FILES["backend/app/reports/__init__.py"] = ""

FILES["backend/app/reports/types.py"] = '''from dataclasses import dataclass, field


@dataclass
class Column:
    key: str
    label_tj: str
    label_key: str | None = None
    align: str = "left"
    width: int = 100
    format: str | None = None


@dataclass
class Section:
    title_tj: str
    columns: list
    rows: list
    summary: dict | None = None


@dataclass
class Report:
    key: str
    title_tj: str
    subtitle_tj: str | None = None
    meta: dict = field(default_factory=dict)
    sections: list = field(default_factory=list)

    def add(self, section):
        self.sections.append(section)
        return self


def fmt_money(v, cur="TJS"):
    try:
        return f"{float(v):,.2f} {cur}"
    except (TypeError, ValueError):
        return str(v)


def fmt_int(v):
    try:
        return f"{int(v):,}"
    except (TypeError, ValueError):
        return str(v)


def fmt_pct(v):
    try:
        return f"{float(v):.2f}%"
    except (TypeError, ValueError):
        return str(v)


def format_cell(value, fmt, cur="TJS"):
    if fmt == "money":
        return fmt_money(value, cur)
    if fmt == "int":
        return fmt_int(value)
    if fmt == "float":
        try:
            return f"{float(value):,.2f}"
        except (TypeError, ValueError):
            return str(value)
    if fmt == "pct":
        return fmt_pct(value)
    if value is None:
        return ""
    return str(value)
'''

FILES["backend/app/reports/catalog.py"] = '''CATALOG = [
    {"key": "pnl", "title_tj": "P&L (Foyda/Zarar)",
     "desc_tj": "Tushum, COGS, yalpi va sof foyda", "note_tj": None,
     "formats": ["pdf", "excel", "csv"]},
    {"key": "sales", "title_tj": "Sotuvlar",
     "desc_tj": "Kunlik sotuv va top mahsulot", "note_tj": None,
     "formats": ["pdf", "excel", "csv"]},
    {"key": "inventory", "title_tj": "Ombor zaxirasi",
     "desc_tj": "Mahsulot boyicha zaxira qiymati",
     "note_tj": "Joriy zaxira snapshot", "formats": ["pdf", "excel", "csv"]},
    {"key": "dead_stock", "title_tj": "Qotib qolgan zaxira",
     "desc_tj": "DEAD / SLOW / AT-RISK",
     "note_tj": "Joriy zaxira snapshot", "formats": ["pdf", "excel", "csv"]},
    {"key": "supplier", "title_tj": "Supplier anomaliyalari",
     "desc_tj": "Narx oshib ketgan supplierlar",
     "note_tj": "Aniqlash oynasi: oxirgi 180 kun",
     "formats": ["pdf", "excel", "csv"]},
    {"key": "variance", "title_tj": "Inventarizatsiya farqi",
     "desc_tj": "Fizik sanoq vs tizim",
     "note_tj": "Report davri ishlatiladi",
     "formats": ["pdf", "excel", "csv"]},
    {"key": "forecast", "title_tj": "Prognoz",
     "desc_tj": "Tushum + stockout + excess",
     "note_tj": "Prognoz oynasi: oxirgi 90 kun",
     "formats": ["pdf", "excel", "csv"]},
    {"key": "billing", "title_tj": "Obuna va tolovlar",
     "desc_tj": "Invoicelar va subscription",
     "note_tj": "Report davri ishlatiladi",
     "formats": ["pdf", "excel", "csv"]},
]
'''

FILES["backend/app/reports/api.py"] = '''from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.deps import require
from app.core.validators import ensure_branch
from app.models.user import User
from app.reports.catalog import CATALOG
from app.reports.builders import BUILDERS
from app.reports.exporters import EXPORTERS, MIME_TYPES, FILE_EXTS

router = APIRouter(prefix="/api/reports", tags=["reports"])


@router.get("/catalog")
def catalog():
    return CATALOG


def _period(since_str, until_str):
    if until_str and len(until_str) == 10:
        until_dt = datetime.fromisoformat(until_str) + timedelta(days=1)
    elif until_str:
        until_dt = datetime.fromisoformat(until_str)
    else:
        until_dt = datetime.utcnow() + timedelta(seconds=1)
    if since_str:
        since_dt = datetime.fromisoformat(since_str)
    else:
        since_dt = until_dt - timedelta(days=30)
    return since_dt, until_dt


@router.get("/{report_key}.{fmt}")
def download(report_key: str, fmt: str, request: Request,
             since: str | None = Query(default=None),
             until: str | None = Query(default=None),
             branch_id: int | None = None,
             user: User = Depends(require("reports", "view")),
             db: Session = Depends(get_db)):
    if report_key not in BUILDERS:
        raise HTTPException(404, f"Nomahlum report: {report_key}")
    if fmt not in EXPORTERS:
        raise HTTPException(400, f"Nomahlum format: {fmt}")
    try:
        since_dt, until_dt = _period(since, until)
    except ValueError:
        raise HTTPException(400, "Notogri sana formati")
    if branch_id is not None:
        ensure_branch(db, user.tenant_id, branch_id)
    builder = BUILDERS[report_key]
    report = builder(db, user.tenant_id, since_dt, until_dt, branch_id=branch_id)
    report.meta["until_exclusive"] = True
    locale = "tj"
    try:
        from app.i18n.resolver import resolve_locale
        locale = resolve_locale(request)
    except Exception:
        pass
    exporter = EXPORTERS[fmt]
    try:
        content = exporter(report, locale=locale)
    except TypeError:
        content = exporter(report)
    fn = f"{report_key}_{since_dt.date()}_{until_dt.date()}.{FILE_EXTS[fmt]}"
    return Response(content=content, media_type=MIME_TYPES[fmt],
                    headers={"Content-Disposition": f'attachment; filename="{fn}"',
                             "Cache-Control": "no-store"})
'''

FILES["backend/app/reports/builders/__init__.py"] = '''from app.reports.builders.pnl import build_pnl
from app.reports.builders.sales import build_sales
from app.reports.builders.inventory import build_inventory
from app.reports.builders.dead_stock import build_dead_stock
from app.reports.builders.supplier import build_supplier
from app.reports.builders.variance import build_variance
from app.reports.builders.forecast import build_forecast
from app.reports.builders.billing import build_billing

BUILDERS = {"pnl": build_pnl, "sales": build_sales, "inventory": build_inventory,
            "dead_stock": build_dead_stock, "supplier": build_supplier,
            "variance": build_variance, "forecast": build_forecast,
            "billing": build_billing}
'''

FILES["backend/app/reports/builders/pnl.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.profit import real_profit, profit_by_branch


def build_pnl(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    p = real_profit(db, tenant_id, since, until, branch_id=branch_id)
    rpt = Report(key="pnl", title_tj="Foyda va zarar hisoboti (P&L)",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    rpt.add(Section(title_tj="Moliyaviy korsatkichlar",
                    columns=[Column("item", "Korsatkich", width=220),
                             Column("value", "Summa", align="right", width=160, format="money")],
                    rows=[{"item": "Tushum", "value": p["revenue"]},
                          {"item": "COGS", "value": p["cogs"]},
                          {"item": "Yalpi foyda", "value": p["gross_profit"]},
                          {"item": "Xarajat", "value": p["expenses"]},
                          {"item": "Sof foyda", "value": p["net_profit"]}]))
    rpt.add(Section(title_tj="Marja va KPI",
                    columns=[Column("item", "Korsatkich", width=220),
                             Column("value", "Qiymat", align="right")],
                    rows=[{"item": "Yalpi marja", "value": f"{p['gross_margin_pct']:.2f}%"},
                          {"item": "Sof marja", "value": f"{p['net_margin_pct']:.2f}%"},
                          {"item": "Sotuvlar soni", "value": f"{p['sales_count']:,}"}]))
    branches = profit_by_branch(db, tenant_id, since, until)
    if len(branches) > 1:
        rpt.add(Section(title_tj="Filiallar boyicha",
                        columns=[Column("name", "Filial"),
                                 Column("revenue", "Tushum", align="right", format="money"),
                                 Column("gross_profit", "Yalpi", align="right", format="money"),
                                 Column("net_profit", "Sof", align="right", format="money")],
                        rows=[{"name": b["branch_name"], "revenue": b["revenue"],
                               "gross_profit": b["gross_profit"],
                               "net_profit": b["net_profit"]} for b in branches]))
    return rpt
'''

FILES["backend/app/reports/builders/sales.py"] = '''from datetime import datetime
from sqlalchemy import func
from app.models.tenant import Tenant
from app.models.sale import Sale, SaleItem
from app.models.product import Product
from app.reports.types import Report, Section, Column


def build_sales(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    q = (db.query(func.date(Sale.created_at).label("day"),
                  func.count(Sale.id).label("count"),
                  func.coalesce(func.sum(Sale.total), 0).label("revenue"),
                  func.coalesce(func.sum(Sale.discount), 0).label("discount"))
         .filter(Sale.tenant_id == tenant_id, Sale.created_at >= since,
                 Sale.created_at < until))
    if branch_id:
        q = q.filter(Sale.branch_id == branch_id)
    rows = q.group_by(func.date(Sale.created_at)).order_by(func.date(Sale.created_at).desc()).all()
    rpt = Report(key="sales", title_tj="Sotuv hisoboti",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    rpt.add(Section(title_tj="Kunlik sotuvlar",
                    columns=[Column("day", "Sana"),
                             Column("count", "Sotuv", align="right", format="int"),
                             Column("revenue", "Tushum", align="right", format="money"),
                             Column("discount", "Skidka", align="right", format="money")],
                    rows=[{"day": str(r.day), "count": r.count,
                           "revenue": float(r.revenue), "discount": float(r.discount)} for r in rows]))
    return rpt
'''

FILES["backend/app/reports/builders/inventory.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.costing import products_valuation, stock_valuation, CostingMethod


def build_inventory(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    products = products_valuation(db, tenant_id, branch_id=branch_id,
                                  method=CostingMethod.MOVING_AVERAGE)
    total = stock_valuation(db, tenant_id, branch_id=branch_id,
                            method=CostingMethod.MOVING_AVERAGE)
    rpt = Report(key="inventory", title_tj="Ombor zaxirasi",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    rpt.add(Section(title_tj="Umumiy",
                    columns=[Column("item", "Korsatkich"),
                             Column("value", "Qiymat", align="right")],
                    rows=[{"item": "Mahsulot turi", "value": f"{len(products):,}"},
                          {"item": "Umumiy miqdor", "value": f"{total['total_qty']:,.2f}"},
                          {"item": "Umumiy qiymat",
                           "value": f"{total['total_value']:,.2f} {cur}"}]))
    rpt.add(Section(title_tj="Mahsulotlar",
                    columns=[Column("sku", "SKU"), Column("name", "Mahsulot"),
                             Column("quantity", "Miqdor", align="right", format="float"),
                             Column("unit_cost", "Tan narx", align="right", format="money"),
                             Column("total_value", "Qiymat", align="right", format="money")],
                    rows=products))
    return rpt
'''

FILES["backend/app/reports/builders/dead_stock.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.dead_stock_tiers import classify, TiersConfig


def build_dead_stock(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    r = classify(db, tenant_id, branch_id=branch_id, cfg=TiersConfig(),
                 include_normal=False)
    rpt = Report(key="dead_stock", title_tj="Qotib qolgan zaxira",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    s = r["summary"]
    rpt.add(Section(title_tj="Umumiy",
                    columns=[Column("tier", "Kategoriya"),
                             Column("count", "Mahsulot", align="right", format="int"),
                             Column("capital_locked", "Kapital", align="right", format="money")],
                    rows=[{"tier": t, "count": s["by_tier"][t]["count"],
                           "capital_locked": s["by_tier"][t]["capital_locked"]}
                          for t in ("DEAD", "SLOW", "AT-RISK", "NORMAL")]))
    for tier in ("DEAD", "SLOW"):
        items = [it for it in r["items"] if it["tier"] == tier]
        if not items:
            continue
        rpt.add(Section(title_tj=f"{tier} mahsulotlar",
                        columns=[Column("sku", "SKU"), Column("name", "Mahsulot"),
                                 Column("quantity", "Miqdor", align="right", format="float"),
                                 Column("capital_locked", "Kapital", align="right", format="money"),
                                 Column("last_sale_days", "Oxirgi sotuv (kun)",
                                        align="right", format="int")],
                        rows=items))
    return rpt
'''

FILES["backend/app/reports/builders/supplier.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.supplier_anomaly import analyze, AnomalyConfig


def build_supplier(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    cfg = AnomalyConfig()
    r = analyze(db, tenant_id, cfg)
    rpt = Report(key="supplier", title_tj="Supplier narx anomaliyalari",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    s = r["summary"]
    rpt.add(Section(title_tj=f"Umumiy (oyna: {cfg.window_days} kun)",
                    columns=[Column("item", "Korsatkich"),
                             Column("value", "Qiymat", align="right")],
                    rows=[{"item": "Tekshirilgan", "value": f"{s['products_analyzed']:,}"},
                          {"item": "Anomaliya", "value": f"{s['products_with_anomalies']:,}"},
                          {"item": "Supplierlar", "value": f"{s['suppliers_flagged']:,}"},
                          {"item": "Umumiy tasir",
                           "value": f"{s['total_estimated_impact']:,.2f} {cur}"}]))
    if r["suppliers"]:
        rpt.add(Section(title_tj="Supplier reytingi",
                        columns=[Column("supplier_name", "Supplier", width=200),
                                 Column("anomaly_count", "Anomaliya",
                                        align="right", format="int"),
                                 Column("estimated_impact", "Tasir",
                                        align="right", format="money")],
                        rows=[{"supplier_name": x["supplier_name"] or f"#{x['supplier_id']}",
                               "anomaly_count": x["anomaly_count"],
                               "estimated_impact": x["estimated_impact"]}
                              for x in r["suppliers"]]))
    return rpt
'''

FILES["backend/app/reports/builders/variance.py"] = '''from datetime import datetime, timedelta
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.variance import analyze, VarianceConfig


def build_variance(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    r = analyze(db, tenant_id, since, until - timedelta(microseconds=1),
                branch_id=branch_id, cfg=VarianceConfig())
    rpt = Report(key="variance", title_tj="Inventarizatsiya farqi",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    s = r["summary"]
    rpt.add(Section(title_tj="Umumiy",
                    columns=[Column("item", "Korsatkich"),
                             Column("value", "Qiymat", align="right")],
                    rows=[{"item": "Sanoq", "value": f"{s['total_counts']:,}"},
                          {"item": "Yoqotish (qiymat)",
                           "value": f"{s['loss_value']:,.2f} {cur}"},
                          {"item": "Ortiqcha (qiymat)",
                           "value": f"{s['surplus_value']:,.2f} {cur}"},
                          {"item": "Sof yoqotish",
                           "value": f"{s['net_loss_value']:,.2f} {cur}"},
                          {"item": "Recurring",
                           "value": f"{s['recurring_offenders']:,}"}]))
    if r["recurring_offenders"]:
        rpt.add(Section(title_tj="Takrorlanuvchi yoqotishlar",
                        columns=[Column("name", "Mahsulot", width=200),
                                 Column("branch_name", "Filial"),
                                 Column("events", "Marta", align="right", format="int"),
                                 Column("total_lost_value", "Qiymat",
                                        align="right", format="money")],
                        rows=r["recurring_offenders"]))
    return rpt
'''

FILES["backend/app/reports/builders/forecast.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.intelligence.forecast import (revenue_forecast, stockout_risk,
                                        excess_stock_risk, ForecastConfig)


def build_forecast(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    cfg = ForecastConfig()
    rev = revenue_forecast(db, tenant_id, branch_id, cfg)
    stockout = stockout_risk(db, tenant_id, branch_id, None, cfg) if branch_id else []
    excess = excess_stock_risk(db, tenant_id, branch_id, None, cfg) if branch_id else []
    rpt = Report(key="forecast", title_tj="Prognoz hisoboti",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    if rev.get("available"):
        rpt.add(Section(title_tj=f"Tushum prognozi (tarix: {cfg.history_days} kun)",
                        columns=[Column("horizon", "Davr"),
                                 Column("total", "Umumiy", align="right", format="money")],
                        rows=[{"horizon": "7 kun", "total": rev["forecast"]["7d"]["total"]},
                              {"horizon": "14 kun", "total": rev["forecast"]["14d"]["total"]},
                              {"horizon": "30 kun", "total": rev["forecast"]["30d"]["total"]}]))
    if stockout:
        rpt.add(Section(title_tj="Stockout xavfi",
                        columns=[Column("name", "Mahsulot", width=200),
                                 Column("current_stock", "Zaxira", align="right", format="float"),
                                 Column("days_until_stockout", "Kungacha",
                                        align="right", format="float"),
                                 Column("risk", "Xavf")],
                        rows=stockout[:30]))
    if excess:
        rpt.add(Section(title_tj="Ortiqcha zaxira",
                        columns=[Column("name", "Mahsulot", width=200),
                                 Column("days_of_cover", "Qoplash (kun)",
                                        align="right", format="float"),
                                 Column("capital_locked", "Kapital",
                                        align="right", format="money"),
                                 Column("risk", "Xavf")],
                        rows=excess[:30]))
    return rpt
'''

FILES["backend/app/reports/builders/billing.py"] = '''from datetime import datetime
from app.models.tenant import Tenant
from app.reports.types import Report, Section, Column
from app.billing.models import Invoice, Subscription
from app.billing.plans import PLANS


def build_billing(db, tenant_id, since, until, branch_id=None):
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    cur = tenant.currency if tenant else "TJS"
    sub = (db.query(Subscription).filter(Subscription.tenant_id == tenant_id)
           .order_by(Subscription.id.desc()).first())
    invoices = (db.query(Invoice).filter(Invoice.tenant_id == tenant_id,
                                          Invoice.issued_at >= since,
                                          Invoice.issued_at < until)
                .order_by(Invoice.issued_at.desc()).limit(500).all())
    rpt = Report(key="billing", title_tj="Obuna va tolovlar",
                 subtitle_tj=tenant.name if tenant else None,
                 meta={"currency": cur, "generated_at": datetime.utcnow().isoformat()})
    if sub:
        plan = PLANS.get(sub.plan_code)
        rpt.add(Section(title_tj="Joriy obuna",
                        columns=[Column("item", "Korsatkich"),
                                 Column("value", "Qiymat", align="right")],
                        rows=[{"item": "Plan",
                               "value": plan.name_tj if plan else sub.plan_code},
                              {"item": "Holat", "value": sub.status}]))
    rpt.add(Section(title_tj="Invoicelar",
                    columns=[Column("number", "Raqam"),
                             Column("plan_code", "Plan"),
                             Column("amount", "Summa", align="right", format="money"),
                             Column("status", "Holat")],
                    rows=[{"number": i.number, "plan_code": i.plan_code,
                           "amount": i.amount, "status": i.status} for i in invoices]))
    return rpt
'''

FILES["backend/app/reports/exporters/__init__.py"] = '''from app.reports.exporters.csv_exporter import to_csv
from app.reports.exporters.excel_exporter import to_excel
from app.reports.exporters.pdf_exporter import to_pdf

EXPORTERS = {"csv": to_csv, "excel": to_excel, "xlsx": to_excel, "pdf": to_pdf}
MIME_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
FILE_EXTS = {"csv": "csv", "excel": "xlsx", "xlsx": "xlsx", "pdf": "pdf"}
'''

FILES["backend/app/reports/exporters/csv_exporter.py"] = '''import csv
import io
from app.reports.types import format_cell


def to_csv(report, locale="tj"):
    buf = io.StringIO()
    w = csv.writer(buf)
    cur = report.meta.get("currency", "TJS")
    w.writerow([report.title_tj])
    if report.subtitle_tj:
        w.writerow([report.subtitle_tj])
    w.writerow([])
    for s in report.sections:
        w.writerow([s.title_tj])
        w.writerow([c.label_tj for c in s.columns])
        for row in s.rows:
            w.writerow([format_cell(row.get(c.key), c.format, cur) for c in s.columns])
        if s.summary:
            w.writerow([format_cell(s.summary.get(c.key, ""), c.format, cur)
                        if c.key in s.summary else "" for c in s.columns])
        w.writerow([])
    return buf.getvalue().encode("utf-8-sig")
'''

FILES["backend/app/reports/exporters/excel_exporter.py"] = '''import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from app.reports.types import format_cell

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F8CFF")
TITLE_FONT = Font(bold=True, size=14)


def to_excel(report, locale="tj"):
    wb = Workbook()
    wb.remove(wb.active)
    cur = report.meta.get("currency", "TJS")
    used = set()
    for idx, s in enumerate(report.sections):
        name = _safe(s.title_tj, used, idx)
        ws = wb.create_sheet(name)
        ws["A1"] = report.title_tj
        ws["A1"].font = TITLE_FONT
        if report.subtitle_tj:
            ws["A2"] = report.subtitle_tj
        ws.append([])
        hr = ws.max_row + 1
        for i, c in enumerate(s.columns, 1):
            cell = ws.cell(row=hr, column=i, value=c.label_tj)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for row in s.rows:
            ws.append([format_cell(row.get(c.key), c.format, cur) for c in s.columns])
        if s.summary:
            ws.append([format_cell(s.summary.get(c.key, ""), c.format, cur)
                       if c.key in s.summary else "" for c in s.columns])
        for i, c in enumerate(s.columns, 1):
            letter = ws.cell(row=hr, column=i).column_letter
            ws.column_dimensions[letter].width = max(12, c.width / 7)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _safe(name, used, idx):
    clean = "".join(ch for ch in (name or "") if ch not in "[]:*?/\\\\")[:28]
    clean = clean or f"Sheet{idx+1}"
    base = clean
    n = 1
    while clean in used:
        clean = f"{base[:26]}_{n}"
        n += 1
    used.add(clean)
    return clean
'''

FILES["backend/app/reports/exporters/pdf_exporter.py"] = '''import io
from pathlib import Path
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from app.reports.types import format_cell

_ASSETS = Path(__file__).parent.parent / "assets"
_REG = _ASSETS / "DejaVuSans.ttf"
_BLD = _ASSETS / "DejaVuSans-Bold.ttf"
_REGISTERED = False
FN, FB = "Helvetica", "Helvetica-Bold"


def _register():
    global _REGISTERED, FN, FB
    if _REGISTERED:
        return
    try:
        if _REG.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(_REG)))
            FN = "DejaVuSans"
        if _BLD.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(_BLD)))
            FB = "DejaVuSans-Bold"
        elif _REG.exists():
            FB = "DejaVuSans"
    except Exception:
        pass
    _REGISTERED = True


_register()
STYLES = getSampleStyleSheet()
TITLE = ParagraphStyle("T", parent=STYLES["Heading1"], fontName=FB,
                       fontSize=16, textColor=colors.HexColor("#1b2029"))
SUB = ParagraphStyle("S", parent=STYLES["Normal"], fontName=FN, fontSize=10,
                     textColor=colors.HexColor("#555555"))
SEC = ParagraphStyle("SEC", parent=STYLES["Heading2"], fontName=FB,
                     fontSize=12, textColor=colors.HexColor("#4f8cff"),
                     spaceBefore=10, spaceAfter=6)


def to_pdf(report, locale="tj"):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    cur = report.meta.get("currency", "TJS")
    el = [Paragraph(_esc(report.title_tj), TITLE)]
    if report.subtitle_tj:
        el.append(Paragraph(_esc(report.subtitle_tj), SUB))
    el.append(Paragraph(_esc(f"Yaratilgan: {report.meta.get('generated_at', '')}"), SUB))
    el.append(Spacer(1, 8))
    for s in report.sections:
        el.append(Paragraph(_esc(s.title_tj), SEC))
        data = [[c.label_tj for c in s.columns]]
        for row in s.rows:
            data.append([format_cell(row.get(c.key), c.format, cur) for c in s.columns])
        if s.summary:
            data.append([format_cell(s.summary.get(c.key, ""), c.format, cur)
                         if c.key in s.summary else "" for c in s.columns])
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4f8cff")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), FB),
            ("FONTNAME", (0,1), (-1,-1), FN),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#cccccc")),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0,1), (-1,-1),
             [colors.white, colors.HexColor("#f5f7fb")]),
        ]))
        el.append(t)
        el.append(Spacer(1, 8))
    doc.build(el)
    return buf.getvalue()


def _esc(s):
    s = "" if s is None else str(s)
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
'''


def main():
    count = 0
    for path, content in FILES.items():
        d = os.path.dirname(path)
        if d:
            os.makedirs(d, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        count += 1
        print(f"OK {path}")
    print(f"\nPART B1: {count} ta fayl")
    subprocess.run(["git", "add", "-A"], check=False)
    subprocess.run(["git", "commit", "-m", "PART B1: Reports"], check=False)
    subprocess.run(["git", "push", "origin", "main"], check=False)
    print("GitHub'ga yuborildi")


if __name__ == "__main__":
    main()
