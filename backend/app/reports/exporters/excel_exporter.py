import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from app.reports.types import format_cell

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F8CFF")
TITLE_FONT = Font(bold=True, size=14)


def to_excel(report, locale="tj"):
    wb = Workbook()
    wb.remove(wb.active)
    cur = report.meta.get("currency", "TJS")
    used = set()
    for idx, s in enumerate(report.sections):
        name = _safe(s.title_tj, used, idx)
        ws = wb.create_sheet(name)
        ws["A1"] = report.title_tj
        ws["A1"].font = TITLE_FONT
        if report.subtitle_tj:
            ws["A2"] = report.subtitle_tj
        ws.append([])
        hr = ws.max_row + 1
        for i, c in enumerate(s.columns, 1):
            cell = ws.cell(row=hr, column=i, value=c.label_tj)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for row in s.rows:
            ws.append([format_cell(row.get(c.key), c.format, cur) for c in s.columns])
        if s.summary:
            ws.append([format_cell(s.summary.get(c.key, ""), c.format, cur)
                       if c.key in s.summary else "" for c in s.columns])
        for i, c in enumerate(s.columns, 1):
            letter = ws.cell(row=hr, column=i).column_letter
            ws.column_dimensions[letter].width = max(12, c.width / 7)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _safe(name, used, idx):
    clean = "".join(ch for ch in (name or "") if ch not in "[]:*?/\\")[:28]
    clean = clean or f"Sheet{idx+1}"
    base = clean
    n = 1
    while clean in used:
        clean = f"{base[:26]}_{n}"
        n += 1
    used.add(clean)
    return clean
